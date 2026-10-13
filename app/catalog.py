"""Validate local Xid rules and replace the active catalog atomically."""
from __future__ import annotations
import copy, hashlib, io, json, os, re, threading, uuid, zipfile
from datetime import datetime, timezone
from pathlib import Path

MAX_UPLOAD=10*1024*1024
DEFAULT={'file':'official-reference-links','sha256':'','codes':{}}


def text(value,field,limit=50000):
    if not isinstance(value,str) or len(value)>limit:raise ValueError(f'{field}格式不正确')
    return value


def basename(value):return value.replace('\\','/').rsplit('/',1)[-1]


def validate(data,allow_empty=False):
    if not isinstance(data,dict) or not isinstance(data.get('codes'),dict):raise ValueError('JSON 缺少 codes 规则表')
    if not (0 if allow_empty else 1)<=len(data['codes'])<=4096:raise ValueError('规则数量须为 1–4096 条')
    file=basename(text(data.get('file',''),'来源名称',512))
    digest=text(data.get('sha256',''),'SHA-256',64)
    if digest and not re.fullmatch('[a-fA-F0-9]{64}',digest):raise ValueError('SHA-256 格式不正确')
    codes={}
    for key,item in data['codes'].items():
        if not isinstance(key,str) or not re.fullmatch(r'0|[1-9][0-9]{0,4}',key) or not isinstance(item,dict):raise ValueError('Xid 编号或规则格式不正确')
        row=item.get('row')
        if row is not None and (type(row) is not int or row<1):raise ValueError(f'Xid {key} 来源行号不正确')
        actions=item.get('actions',[])
        if not isinstance(actions,list) or len(actions)>16:raise ValueError(f'Xid {key} 建议格式不正确')
        parsed=[]
        for action in actions:
            if not isinstance(action,dict):raise ValueError(f'Xid {key} 建议格式不正确')
            parsed.append({name:text(action.get(name,''),'建议') for name in ('key','label','text')})
        codes[key]={'title':text(item.get('title',''),'标题'), 'description':text(item.get('description',''),'说明'),
                    'actions':parsed,'symbol':text(item.get('symbol',''),'符号'),
                    'sheet':text(item.get('sheet',''),'工作表',256),'row':row,
                    'association154':text(item.get('association154',''),'关联说明')}
        architectures=item.get('architectures',{})
        if not isinstance(architectures,dict) or len(architectures)>100 or any(not isinstance(k,str) or not isinstance(v,(str,int,float,bool,type(None))) for k,v in architectures.items()):raise ValueError('架构字段格式不正确')
        codes[key]['architectures']=architectures
    result={'file':file,'sha256':digest.lower(),'codes':codes}
    if data.get('sourceSha256'):result['sourceSha256']=text(data['sourceSha256'],'来源 SHA-256',64)
    if data.get('importFile'):result['importFile']=basename(text(data['importFile'],'导入文件',512))
    return result


def unique_pairs(pairs):
    result={}
    for key,value in pairs:
        if key in result:raise ValueError('JSON 包含重复字段或 Xid 编号')
        result[key]=value
    return result


def read_json(raw):
    try:return json.loads(raw.decode('utf-8-sig'),object_pairs_hook=unique_pairs,parse_constant=lambda _: (_ for _ in ()).throw(ValueError('JSON 含无效数值')))
    except (UnicodeError,RecursionError,json.JSONDecodeError) as exc:raise ValueError('JSON 文件格式不正确') from exc


def workbook(raw):
    import openpyxl
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members=archive.infolist()
            if len(members)>1000 or sum(m.file_size for m in members)>64*1024*1024:raise ValueError('表格内容过大')
            for member in members:
                if member.filename.endswith('.xml'):
                    content=archive.read(member).replace(b'\x00',b'').upper()
                    if b'<!DOCTYPE' in content or b'<!ENTITY' in content:raise ValueError('表格含不支持的 XML 内容')
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            if not {'Xids','Resolution Buckets'}<=set(book.sheetnames):raise ValueError('表格需包含 Xids 和 Resolution Buckets 工作表')
            if (book['Xids'].max_column or 0)<12:raise ValueError('Xids 工作表需保留原表的 12 列')
            for name in ('Xids','Resolution Buckets'):
                if (book[name].max_row or 0)>10000:raise ValueError('工作表行数过多')
            buckets={}
            for row in book['Resolution Buckets'].iter_rows(min_row=2,max_col=2,values_only=True):
                if row[0]:buckets[str(row[0]).split('\n')[0]]={'label':str(row[0]),'text':str(row[1] or '')}
            codes={}
            for n,row in enumerate(book['Xids'].iter_rows(min_row=2,max_col=12,values_only=True),2):
                value=row[1]
                if value is None:continue
                if isinstance(value,bool) or not isinstance(value,(int,float,str)):raise ValueError(f'第 {n} 行 Xid 编号不正确')
                try:
                    code=str(int(value))
                    if isinstance(value,float) and value!=int(value):raise ValueError()
                    if isinstance(value,str) and value.strip()!=code:raise ValueError()
                except (ValueError,OverflowError):raise ValueError(f'第 {n} 行 Xid 编号不正确')
                if code in codes:raise ValueError(f'Xid {code} 重复')
                actions=[]
                for value in row[8:10]:
                    if value:
                        key=str(value).split('\n')[0]
                        actions.append({'key':key,'label':str(value),'text':buckets.get(key,{}).get('text','')})
                codes[code]={'title':str(row[3] or row[2] or ''),'symbol':str(row[2] or ''),
                            'description':str(row[11] or ''),'actions':actions,'sheet':'Xids','row':n,
                            'architectures':dict(zip(['A100','H100','B100','GB200'],row[4:8])),
                            'association154':str(row[10] or '')}
            return {'file':'','sha256':'','codes':codes}
        finally:book.close()
    except ValueError:raise
    except Exception as exc:raise ValueError('XLSX 文件无法读取') from exc


def parse_upload(raw,name):
    if not 0<len(raw)<=MAX_UPLOAD:raise ValueError('请选择 10 MB 以内的规则文件')
    name=basename(name)
    if len(name)>255 or not name or '\x00' in name:raise ValueError('文件名不正确')
    suffix=Path(name).suffix.lower()
    if suffix=='.xlsx':data=workbook(raw);data['file']=name
    elif suffix=='.json':data=read_json(raw)
    else:raise ValueError('请选择 XLSX 或 JSON 文件')
    data=validate(data)
    if data['sha256']:data['sourceSha256']=data['sha256']
    data['sha256']=hashlib.sha256(raw).hexdigest();data['importFile']=name
    if not data['file']:data['file']=name
    return data


def initial_catalog(root,app_home):
    chosen=Path(os.environ.get('XIDLENS_CATALOG',str(app_home/'xid_catalog.local.json' if (app_home/'xid_catalog.local.json').is_file() else root/'xid_catalog.json')))
    try:
        data=validate(read_json(chosen.read_bytes()),allow_empty=True)
        if data['codes']:data=parse_upload(chosen.read_bytes(),chosen.name)
        return data,''
    except (OSError,ValueError):return copy.deepcopy(DEFAULT),'启动规则无法读取，当前使用默认规则。'


class CatalogStore:
    def __init__(self,runtime,initial=None,warning=''):
        self.path=Path(runtime)/'rules'/'active.json';self.lock=threading.RLock()
        self.data=copy.deepcopy(initial if initial is not None else DEFAULT);self.imported_at='';self.warning=warning
        if self.path.exists():
            try:
                saved=read_json(self.path.read_bytes());data=validate(saved['catalog'],allow_empty=True)
                stamp=text(saved.get('importedAt',''),'导入时间',64)
                self.data=data;self.imported_at=stamp;self.warning=''
            except (OSError,ValueError,KeyError,TypeError):self.warning='保存的规则无法读取，当前使用启动规则。'
    def snapshot(self):
        with self.lock:return copy.deepcopy(self.data)
    def info(self):
        with self.lock:
            return {k:v for k,v in self.data.items() if k!='codes'}|{'count':len(self.data['codes']),'importedAt':self.imported_at,'warning':self.warning}
    def apply(self,data):
        with self.lock:
            stamp=datetime.now(timezone.utc).isoformat(timespec='seconds')
            self.path.parent.mkdir(parents=True,exist_ok=True)
            tmp=self.path.parent/(uuid.uuid4().hex+'.pending')
            try:
                with tmp.open('w',encoding='utf-8') as f:
                    json.dump({'catalog':data,'importedAt':stamp},f,ensure_ascii=False);f.flush();os.fsync(f.fileno())
                tmp.replace(self.path)
            finally:tmp.unlink(missing_ok=True)
            self.data=copy.deepcopy(data);self.imported_at=stamp;self.warning=''
            return self.info()
    def import_file(self,raw,name):return self.apply(parse_upload(raw,name))
    def reset(self):return self.apply(copy.deepcopy(DEFAULT))
